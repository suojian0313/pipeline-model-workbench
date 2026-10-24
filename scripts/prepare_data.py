"""只读试题附件，生成两种后端共用的数据；可从这里学习 Python 文件与集合操作。"""
import argparse
import csv
import json
import re
import struct
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def normalize_code(value):
    """仅对业务标识清理空白，原始属性仍保留给界面展示。"""
    result = re.sub(r"\s+", "", str(value if value is not None else ""))
    return "" if result == "-" else result


def iso_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ValueError(f"焊接日期必须为 YYYY-MM-DD 或空值：{value!r}")
    return date.fromisoformat(value).isoformat()


def read_b3dm(path):
    data = Path(path).read_bytes()
    magic, version, length, ft_json, ft_bin, bt_json, _ = struct.unpack_from("<4s6I", data)
    if magic != b"b3dm" or version != 1 or length != len(data):
        raise ValueError(f"不支持或损坏的 b3dm：{path}")
    start = 28 + ft_json + ft_bin
    return json.loads(data[start:start + bt_json].decode("utf-8").rstrip("\x00 "))


def merge_bounds(previous, minimum, maximum):
    if previous is None:
        return {"min": list(minimum), "max": list(maximum)}
    return {
        "min": [min(a, b) for a, b in zip(previous["min"], minimum)],
        "max": [max(a, b) for a, b in zip(previous["max"], maximum)],
    }


def build_model(structure, rows, batches, asset_dir):
    by_path = {row["tree_key"]: row for row in rows}
    if len(by_path) != len(rows):
        raise ValueError("CSV tree_key 重复，不能自动选择其中一行")
    geometry_bounds = {}
    for batch in batches:
        ids = batch["componentId"]
        if len(batch["boundsMin"]) != 3 * len(ids) or len(batch["boundsMax"]) != 3 * len(ids):
            raise ValueError("Batch Table bounds 数量与 componentId 不匹配")
        for i, geometry_id in enumerate(ids):
            minimum = batch["boundsMin"][i * 3:i * 3 + 3]
            maximum = batch["boundsMax"][i * 3:i * 3 + 3]
            geometry_bounds[geometry_id] = merge_bounds(geometry_bounds.get(geometry_id), minimum, maximum)

    components, geometry_to_component, node_to_component = [], {}, {}
    visited_paths, original_geometry = set(), set()
    tree_nodes = 0

    def visit(node, depth, path, owner=None):
        nonlocal tree_nodes
        tree_nodes += 1
        if depth == 5:
            if path not in by_path:
                raise ValueError(f"结构树业务节点找不到属性：{path}")
            row = by_path[path]
            visited_paths.add(path)
            owner = {
                "id": node["id"], "name": node["name"],
                "line_no": normalize_code(row.get("EP3D_管 线 号")),
                "kind": row["component_name"],
                "properties": {key: value for key, value in row.items() if key.startswith("EP3D_")},
                "geometry_ids": [], "bounds": None,
                "weld_no": normalize_code(row.get("EP3D_焊点编号")),
                "material_code": normalize_code(row.get("EP3D_元件编号")),
                "material_1": normalize_code(row.get("EP3D_元件1编号")),
                "material_2": normalize_code(row.get("EP3D_元件2编号")),
            }
            components.append(owner)
        geometry_id = node.get("componentId")
        if geometry_id and owner is not None:
            original_geometry.add(geometry_id)
            node_to_component[node["id"]] = owner["id"]
            if geometry_id in geometry_bounds:
                if geometry_id in geometry_to_component and geometry_to_component[geometry_id] != owner["id"]:
                    raise ValueError(f"几何 ID 对应多个业务对象：{geometry_id}")
                geometry_to_component[geometry_id] = owner["id"]
                owner["geometry_ids"].append(geometry_id)
                bounds = geometry_bounds[geometry_id]
                owner["bounds"] = merge_bounds(owner["bounds"], bounds["min"], bounds["max"])
        counts = defaultdict(int)
        for child in node.get("children", []):
            counts[child["name"]] += 1
            segment = child["name"]
            if depth + 1 in (4, 5):
                segment += f"[{counts[child['name']]}]"
            visit(child, depth + 1, path + "/" + segment, owner)

    for root in structure["roots"]:
        visit(root, 0, root["name"])
    if set(by_path) != visited_paths:
        raise ValueError("部分 CSV 属性行没有对应的结构树业务节点")
    if set(geometry_bounds) != set(geometry_to_component):
        raise ValueError("部分 b3dm 构件没有映射到结构树业务对象")

    materials = defaultdict(list)
    for component in components:
        component["geometry_ids"] = sorted(set(component["geometry_ids"]))
        if component["material_code"]:
            materials[(component["line_no"], component["material_code"])].append(component["id"])
    model_welds = []
    for component in components:
        if not component["weld_no"]:
            continue
        sides = [sorted(materials.get((component["line_no"], component[f"material_{side}"]), [])) for side in (1, 2)]
        model_welds.append({
            "line_no": component["line_no"], "weld_no": component["weld_no"],
            "component_id": component["id"], "side_1_ids": sides[0], "side_2_ids": sides[1],
            "missing_sides": [side for side in (1, 2) if not sides[side - 1]],
        })
    unique_keys = {(item["line_no"], item["weld_no"]) for item in model_welds}
    if len(unique_keys) != len(model_welds):
        raise ValueError("模型焊点联合键重复")
    return {
        "source": {"asset_dir": asset_dir},
        "summary": {
            "business_components": len(components), "geometry_components": len(geometry_bounds),
            "model_welds": len(model_welds), "tree_nodes": tree_nodes,
            "original_geometry_nodes": len(original_geometry),
            "unexported_geometry_nodes": len(original_geometry - set(geometry_bounds)),
            "missing_endpoint_refs": sum(len(item["missing_sides"]) for item in model_welds),
        },
        "tree": structure["roots"], "components": components,
        "geometry_to_component": geometry_to_component, "node_to_component": node_to_component,
        "model_welds": model_welds,
    }


def export_excel(path, output, model_keys):
    # read_only 模式逐行读取，避免把 8 万行构造成前端响应。
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    seen, found = set(), set()
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(rows)]
        required = {"管线号", "焊点编号", "材料唯一码1", "材料唯一码2", "焊接日期"}
        if not required.issubset(headers):
            raise ValueError(f"Excel 缺少列：{sorted(required - set(headers))}")
        with output.open("w", encoding="utf-8", newline="\n") as stream:
            for number, values in enumerate(rows, start=2):
                row = dict(zip(headers, values))
                line_no, weld_no = normalize_code(row.get("管线号")), normalize_code(row.get("焊点编号"))
                if not line_no and not weld_no:
                    continue
                key = (line_no, weld_no)
                if not line_no or not weld_no or key in seen:
                    raise ValueError(f"Excel 第 {number} 行焊点键缺失或重复：{key}")
                seen.add(key)
                if key in model_keys:
                    found.add(key)
                item = {"line_no": line_no, "weld_no": weld_no,
                        "material_1": normalize_code(row.get("材料唯一码1")),
                        "material_2": normalize_code(row.get("材料唯一码2")),
                        "welded_on": iso_date(row.get("焊接日期")), "is_model": key in model_keys}
                stream.write(json.dumps(item, ensure_ascii=False) + "\n")
        if model_keys - found:
            raise ValueError(f"模型焊点在 Excel 中不存在：{sorted(model_keys - found)[:5]}")
        return len(seen)
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path)
    args = parser.parse_args()
    source = args.source_dir
    if source is None:
        candidates = [path for path in ROOT.iterdir() if path.is_dir() and list(path.glob("*.xlsx"))]
        if len(candidates) != 1:
            raise ValueError("请用 --source-dir 指定试题附件目录")
        source = candidates[0]
    model_dir = next(path.parent for path in source.rglob("tileset.json"))
    with (model_dir / "全部对象属性.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    structure = json.loads((model_dir / "structure.json").read_text(encoding="utf-8-sig"))
    tiles = sorted(model_dir.glob("*.b3dm"))
    asset_dir = str(model_dir.resolve().relative_to(ROOT)).replace("\\", "/")
    model = build_model(structure, rows, [read_b3dm(path) for path in tiles], asset_dir)
    output_dir = ROOT / "data"
    output_dir.mkdir(exist_ok=True)
    temporary = output_dir / "welds.jsonl.tmp"
    try:
        model_keys = {(item["line_no"], item["weld_no"]) for item in model["model_welds"]}
        count = export_excel(next(source.glob("*.xlsx")), temporary, model_keys)
        model["summary"].update({"tile_files": len(tiles), "source_weld_rows": count,
                                 "tile_bytes": sum(path.stat().st_size for path in tiles)})
        temporary.replace(output_dir / "welds.jsonl")
        (output_dir / "model.json").write_text(json.dumps(model, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    finally:
        temporary.unlink(missing_ok=True)
    print(json.dumps(model["summary"], ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
